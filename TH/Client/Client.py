import socket
from tkinter import *
from tkinter import messagebox
from PIL import Image, ImageTk
import tkinter.ttk as exTk
import tkinter.scrolledtext as scllText
from docx import Document
import re
import os
from fpdf import FPDF
import textwrap

path_download = "Downloads\ID_"

##############################################################
#------------------------ IP CONNECT ------------------------#
def Exit(window):
    window.destroy()

def Send(data):
    s.send(data.encode())

def Connect(host):
    port = 80
    try:
        s.connect((str(host), port))
        messagebox.showinfo("Status", "Successfully")
    except:
        messagebox.showinfo("Status", "Unsuccessfully")


##############################################################
#-------------------------- SIGN IN -------------------------#
def Sub(window, kind, usn, pwd):
    print(kind + "+" + usn + "+" + pwd)
    Send(kind + "+" + usn + "+" + pwd)
    data = s.recv(1024000).decode("utf-8")
    if data == "Success":
        window_sch(window)

    else:
        messagebox.showinfo("Status", "Unsuccessfully")
def SignIn(window):
    window_sgin = Toplevel(window)
    window_sgin.title("Sign in")
    window_sgin.geometry("400x200")
    window_sgin.configure(bg = 'pink')
    img_sgin = Image.open("Image\signin.jpg")
    img_sgin = img_sgin.resize((100, 27), Image.ANTIALIAS)
    imge = ImageTk.PhotoImage(img_sgin)
    label = Label(window_sgin, image = imge)
    label.pack(pady = 10)

    Label_usn = Label(window_sgin, text = "Username")
    Label_usn.place( x = 40, y = 60)

    Label_pwd = Label(window_sgin, text = "Password")
    Label_pwd.place( x = 40, y = 100)

    Entry_usn = Entry(window_sgin, width = 30)
    Entry_usn.place(x = 110, y = 60)

    Entry_pwd = Entry(window_sgin, width = 30, show = '*')
    Entry_pwd.place(x = 110, y = 100)

    btn_submit = Button(window_sgin, text = "Sign in", fg = "navy", bg = "snow", width = 10,
                            command = lambda: Sub(window_sgin, "Signin",Entry_usn.get(), Entry_pwd.get()))
    btn_submit.place(x = 163, y = 140)
        
    window_sgin.mainloop()


##############################################################
#-------------------------- SIGN UP -------------------------#
def Check_pwd(window, usn, pwd1, pwd2):
    if pwd1 == pwd2:
        Sub(window, "Signup", usn, pwd2)
    else:
        messagebox.showinfo("Status", "Password incorrect")
       
def SignUp(window):
    window_sgup = Toplevel(window)
    window_sgup.title("Sign Up")
    window_sgup.geometry("400x280")
    window_sgup.configure(bg = 'pink')
    img_sgup = Image.open("Image\signup.jpg")
    img_sgup = img_sgup.resize((100, 27), Image.ANTIALIAS)
    imge = ImageTk.PhotoImage(img_sgup)
    label = Label(window_sgup, image = imge)
    label.pack(pady = 10)
        
    Label_usn = Label(window_sgup, text = "Username")
    Label_usn.place( x = 40, y = 70)

    Label_pwd = Label(window_sgup, text = "Password")
    Label_pwd.place( x = 40, y = 110)

    Label_pwd2 = Label(window_sgup, text = "Re-enter password")
    Label_pwd2.place( x = 40, y = 150)

    Entry_usn = Entry(window_sgup, width = 30)
    Entry_usn.place(x = 150, y = 70)

    Entry_pwd = Entry(window_sgup, width = 30, show = '*')
    Entry_pwd.place(x = 150, y = 110)

    Entry_pwd2 = Entry(window_sgup, width = 30, show = '*')
    Entry_pwd2.place(x = 150, y = 150)

    btn_submit = Button(window_sgup, text = "Sign up", fg = "navy", bg = "snow", width = 10)
    btn_submit.place(x = 163, y = 200)
    btn_submit['command'] = lambda: Check_pwd(window_sgup, Entry_usn.get(), Entry_pwd.get(), Entry_pwd2.get())
        
    window_sgup.mainloop()


##############################################################
#-------------------------- SEARCH --------------------------#
def Sub2(window_sch, a, b, c, d):
    Send(a + "+" + b + "+" + c + "++" + d)
    data = s.recv(1024000).decode("utf-8")
    if data == "Not found":
        messagebox.showinfo("Status", "Not found!")
    else:
        check_list = data.find("*list\n")
        if check_list == -1:
            (id_book, _data) = data.split("||\n")
            (ID, title_book) = id_book.split('+')
            view_Book(window_sch, _data, ID, title_book)
        else:
            list_Book(window_sch, data)

def window_sch(window):
    window_sch = Toplevel(window)
    window_sch.title("Search")
    window_sch.geometry("400x370")
    window_sch.configure(bg = 'pink')
    img_sch = Image.open("Image\search.jpg")
    img_sch = img_sch.resize((100, 27), Image.ANTIALIAS)
    imge = ImageTk.PhotoImage(img_sch)
    label = Label(window_sch, image = imge)
    label.pack(pady = 15)

    Label_ID = Label(window_sch, text = "F_ID:", font = "Arial 10", width = 7)
    Label_ID.place(x = 20, y = 80)

    Label_Name = Label(window_sch, text = "F_Name:", font = "Arial 10", width = 7)
    Label_Name.place(x = 20, y = 130)

    Label_Type = Label(window_sch, text = "F_Type:", font = "Arial 10", width = 7)
    Label_Type.place(x = 20, y = 180)

    Label_Author = Label(window_sch, text = "F_Author: ", font = "Arial 10", width = 7)
    Label_Author.place(x = 20, y = 230)

    Entry_ID = Entry(window_sch, width = 45)
    Entry_ID.place(x = 95, y = 80)

    Entry_Name = Entry(window_sch, width = 45)
    Entry_Name.place(x = 95, y = 130)

    Entry_Type = Entry(window_sch, width = 45)
    Entry_Type.place(x = 95, y = 180)

    Entry_Author = Entry(window_sch, width = 45)
    Entry_Author.place(x = 95, y = 230)

    btn_sch = Button(window_sch, text = "Search", font = "Arial 10", bg = "limegreen", width = 10)
    btn_sch.place(x = 50, y = 290)
    btn_sch['command'] = lambda: Sub2(window_sch, Entry_ID.get(), Entry_Name.get(), Entry_Author.get(), Entry_Type.get())

    btn_ext = Button(window_sch, text = "Exit", font = "Arial 10", bg = "red", width = 10)
    btn_ext.place(x = 250, y = 290)
    btn_ext['command'] = lambda: Exit(window_sch)

    window_sch.mainloop()


################################################################
#-------------------------- VIEW BOOK -------------------------#
def convert_TO_docx(path, ID, book_title, format):
    document = Document()
    myfile = open(path, 'r')
    data = myfile.read()
    data = re.sub(r'[^\x00-\x7F]+|\x0c',' ', data)
    p = document.add_paragraph(data)
    document.save(path_download + ID + '-' + book_title + format)
    myfile.close()

def convert_TO_pdf(path, ID, book_title, format):
    a4_width_mm = 210
    pt_to_mm = 0.35
    fontsize_pt = 10
    fontsize_mm = fontsize_pt * pt_to_mm
    margin_bottom_mm = 10
    character_width_mm = 7 * pt_to_mm
    width_text = a4_width_mm / character_width_mm

    pdf = FPDF(orientation = 'P', unit = 'mm', format = 'A4')
    pdf.set_auto_page_break(True, margin = margin_bottom_mm)
    pdf.add_page()
    pdf.set_font(family = 'Courier', size = fontsize_pt)

    myfile = open(path, 'r')
    text = myfile.read()
    splitted = text.split('\n')

    for line in splitted:
        lines = textwrap.wrap(line, width_text)

        if len(lines) == 0:
            pdf.ln()

        for wrap in lines:
            pdf.cell(0, fontsize_mm, wrap, ln = 1)

    pdf.output(path_download + ID + '-' + book_title + format, 'F')
    myfile.close()

def download(data, format, ID, book_title): 
    if len(format) != 0:
        filename = "temp.txt"
        f = open(filename, 'w')
        f.write(data)
        f.close()

        if format == ".docx" or format == ".doc":
            convert_TO_docx(filename, ID, book_title, format)
        else:
            if format == ".txt":
                fo = open(path_download + ID + '-' + book_title + format, 'w')
                fi = open(filename, 'r')
                fo.write(fi.read())
                fi.close()
                fo.close()
            else:
                convert_TO_pdf(filename, ID, book_title, format)


def view_Book(window_sch, data, ID, book_title):
    window_view = Toplevel(window_sch)
    window_view.title(ID + '-' + book_title)
    window_view.geometry("1000x500")
    window_view.configure(bg = 'pink')

    editArea = scllText.ScrolledText(window_view, wrap = WORD) 
    editArea.pack(padx = 5, pady = 35, fill = BOTH, expand = True) 
    editArea.insert(INSERT, data)

    combobox = exTk.Combobox(window_view, width = 15, values = ('.txt', '.doc', '.docx', '.pdf'))
    combobox.place(relx = 0.8, rely = 0.945)

    btn_download = Button(window_view, text = "Download", fg = "navy", bg = "snow")
    btn_download.place(relx = 0.93, rely = 0.94)
    btn_download['command'] = lambda:download(data, combobox.get(), ID, book_title)

    window_view.mainloop()

def Sub3(window, data):
    if len(data) != 0:
        (_ID, _name, _author, _pbsh, _Type) = data.split(',')
        Sub2(window, _ID, _name, _author, _Type)

def list_Book(window_sch, data):
    mylist = list(data.split('\n'))
    mylist.pop(-1)
    mylist.pop(0)
    window_select_books = Toplevel(window_sch)
    window_select_books.title("Books List")
    window_select_books.geometry("400x150")
    window_select_books.configure(bg = 'pink')

    label_text = Label(window_select_books, text = "Select the Book : (ID, Name, Author, Publishing year, Type)", fg = "navy")
    label_text.pack(pady = 10)

    combobox = exTk.Combobox(window_select_books, width = 50, values = mylist)
    combobox.pack()
    btn_submit = Button(window_select_books, text = "Submit", fg = "navy", bg = "snow", width = 10)
    btn_submit.pack(pady = 15)
    btn_submit['command'] = lambda:Sub3(window_select_books, combobox.get())

    window_select_books.mainloop()


##################################################################
##################################################################
#-------------------------- MAIN WINDOW -------------------------#

window1 = Tk()
window1.title("Client")
window1.geometry("600x400")
window1.configure(bg = 'pink')
s = socket.socket(socket.AF_INET,socket.SOCK_STREAM)
img_tmp = Image.open("Image\image.jpg")
img_tmp = img_tmp.resize((500, 135), Image.ANTIALIAS)
img = ImageTk.PhotoImage(img_tmp)
panel = Label(window1, image = img)
panel.pack(pady = 10)
Label_IP = Label(window1, text = "IP", fg = "navy", font = "Arial")
Label_IP.place(x = 130, y = 165)

Entry_IP = Entry(window1, bd = 3, bg = "gainsboro", fg = "red", width = 40, text = "Enter IP")
Entry_IP.pack(pady = 10)

btn_cnt = Button(window1, text = "Connect", fg = "navy", bg = "snow", width = 10)
btn_cnt.place(x = 450, y = 165)

btn_cnt["command"] = lambda: Connect(Entry_IP.get())

label_text = Label(window1, text = "Are you sure the connection is successful already?", fg = "navy")
label_text.place(x = 168, y = 250)

btn_sgin = Button(window1, text = "Sign in", fg = "navy", bg = "snow", width = 10)
btn_sgin.place(x = 210, y = 280)

btn_sgup = Button(window1, text = "Sign up", fg = "navy", bg = "snow", width = 10)
btn_sgup.place(x = 310, y = 280)

btn_sgin['command'] = lambda:SignIn(window1)
btn_sgup['command'] = lambda:SignUp(window1)

window1.mainloop()
